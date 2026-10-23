import base64
import io
import json
import os
import re

import streamlit as st
from docx import Document
from docx.shared import Pt
from groq import Groq
from pypdf import PdfReader

# ------------------------------------------------------------------
# Config
# ------------------------------------------------------------------
# Tried in order; if one is retired (404 model_not_found) the next is used.
# Check console.groq.com/docs/models occasionally, since Groq retires models often.
MODELS = ["openai/gpt-oss-120b", "openai/gpt-oss-20b"]

# Vision-capable models for image resumes, tried in order.
# Check console.groq.com/docs/vision and update if you get "model not found".
VISION_MODELS = [
    "qwen/qwen3.8-27b",
    "qwen/qwen3.6-27b",
    "meta-llama/llama-4-scout-17b-16e-instruct",
]

IMAGE_TYPES = ["png", "jpg", "jpeg"]
MAX_IMAGE_BYTES = 4 * 1024 * 1024  # keep images under ~4 MB
MAX_RESUME_CHARS = 12000

st.set_page_config(page_title="Skill-Gap & Career Roadmap Tool", page_icon="🎯", layout="wide")

st.title("🎯 Student Skill-Gap & Career Roadmap Tool")
st.write("Analyze your tech skills, identify gaps for your dream job, and get an AI-powered roadmap.")

# ------------------------------------------------------------------
# Role requirements database (add more roles/skills freely)
# ------------------------------------------------------------------
roles_requirements = {
    "Software Developer": [
        "Python", "JavaScript", "Data Structures & Algorithms", "Databases/SQL",
        "Git", "REST APIs", "Docker", "System Design Basics",
    ],
    "Frontend Developer": [
        "HTML/CSS", "JavaScript", "React", "TypeScript", "Git",
        "Responsive Design", "REST APIs", "Testing",
    ],
    "Backend Developer": [
        "Python/Java/Node.js", "Databases/SQL", "REST APIs", "Authentication",
        "Git", "Docker", "Caching", "System Design Basics",
    ],
    "Data Scientist": [
        "Python", "SQL", "Statistics", "Machine Learning", "Pandas",
        "Data Visualization", "Scikit-learn", "Deep Learning Basics",
    ],
    "Data Analyst": [
        "SQL", "Excel", "Python", "Pandas", "Data Visualization",
        "Power BI/Tableau", "Statistics", "Communication",
    ],
    "Cloud/DevOps Engineer": [
        "Linux", "Docker", "Kubernetes", "AWS/Cloud", "Networking",
        "Terraform", "CI/CD", "Git",
    ],
}


# ------------------------------------------------------------------
# Helpers
# ------------------------------------------------------------------
def checkbox_key(role: str, skill: str) -> str:
    return f"skill_{role}_{skill}"


HARDCODED_API_KEY = ""  # Leave empty. Use Streamlit secrets instead.


def get_api_key() -> str:
    """Key comes from the code, Streamlit secrets, or an environment variable.
    Visitors never see or enter it."""
    if HARDCODED_API_KEY:
        return HARDCODED_API_KEY
    try:
        if "GROQ_API_KEY" in st.secrets:
            return st.secrets["GROQ_API_KEY"]
    except Exception:
        pass
    return os.getenv("GROQ_API_KEY", "")


def ask_llm(prompt: str, temperature: float = 0.5, image_b64=None, mime="image/png"):
    """Call Groq and return text, or None after showing an error.
    If image_b64 is given, a vision model is used."""
    api_key = get_api_key()
    if not api_key:
        st.error("The AI service isn't configured yet. Please try again later.")
        return None

    if image_b64:
        models = VISION_MODELS
        content = [
            {"type": "text", "text": prompt},
            {"type": "image_url", "image_url": {"url": f"data:{mime};base64,{image_b64}"}},
        ]
    else:
        models = MODELS
        content = prompt

    client = Groq(api_key=api_key)
    last_error = None
    for model in models:
        try:
            response = client.chat.completions.create(
                messages=[{"role": "user", "content": content}],
                model=model,
                temperature=temperature,
            )
            return response.choices[0].message.content
        except Exception as e:
            last_error = e
            msg = str(e)
            if "model_not_found" in msg or "404" in msg or "decommissioned" in msg:
                continue  # try the next model
            break  # other errors (bad key, rate limit) won't be fixed by another model
    st.error(f"Groq API error: {last_error}")
    return None


def parse_json_list(text: str):
    """Pull a JSON list out of a model reply, even if it adds extra words."""
    text = re.sub(r"<think>.*?</think>", "", text, flags=re.DOTALL)
    try:
        start, end = text.index("["), text.rindex("]") + 1
        data = json.loads(text[start:end])
        return [str(x) for x in data] if isinstance(data, list) else []
    except (ValueError, json.JSONDecodeError):
        return []


def extract_resume_text(uploaded) -> str:
    """Read text from a PDF, DOCX or TXT upload."""
    name = uploaded.name.lower()
    if name.endswith(".pdf"):
        reader = PdfReader(uploaded)
        return "".join(page.extract_text() or "" for page in reader.pages)
    if name.endswith(".docx"):
        doc = Document(uploaded)
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    parts.append(cell.text)
        return "\n".join(parts)
    if name.endswith(".txt"):
        return uploaded.getvalue().decode("utf-8", errors="ignore")
    return ""


def clean_roadmap_text(text: str) -> str:
    """Safety net: turn <br> into new lines and drop any other HTML tags."""
    text = re.sub(r"<br\s*/?>", "\n", text, flags=re.IGNORECASE)
    return re.sub(r"</?[a-zA-Z][^>]*>", "", text)


def add_runs(paragraph, text: str):
    """Add text to a Word paragraph, turning **bold** and `code` into formatting."""
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            paragraph.add_run(part[1:-1]).font.name = "Consolas"
        else:
            paragraph.add_run(part)


def markdown_to_docx(md_text: str, title: str) -> bytes:
    """Convert the roadmap markdown into a Word document and return its bytes."""
    doc = Document()
    doc.add_heading(title, level=0)
    lines = md_text.splitlines()
    i = 0
    while i < len(lines):
        raw = lines[i].rstrip()
        stripped = raw.strip()

        if not stripped or re.fullmatch(r"[-*_]{3,}", stripped):
            i += 1
            continue

        # Tables (in case the model still produces one)
        if stripped.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            if rows:
                ncols = max(len(r) for r in rows)
                table = doc.add_table(rows=len(rows), cols=ncols)
                table.style = "Table Grid"
                for r_idx, row in enumerate(rows):
                    for c_idx in range(ncols):
                        cell_text = row[c_idx] if c_idx < len(row) else ""
                        add_runs(table.cell(r_idx, c_idx).paragraphs[0], cell_text)
            continue

        # Headings
        m = re.match(r"^(#{1,6})\s+(.*)", stripped)
        if m:
            doc.add_heading(m.group(2).replace("**", ""), level=min(len(m.group(1)), 3))
            i += 1
            continue

        # Bullets
        m = re.match(r"^(\s*)[-*•]\s+(.*)", raw)
        if m:
            style = "List Bullet 2" if len(m.group(1)) >= 2 else "List Bullet"
            add_runs(doc.add_paragraph(style=style), m.group(2))
            i += 1
            continue

        # Numbered lines keep their own numbers
        m = re.match(r"^\s*(\d+[.)])\s+(.*)", raw)
        if m:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(14)
            add_runs(p, f"{m.group(1)} {m.group(2)}")
            i += 1
            continue

        add_runs(doc.add_paragraph(), stripped)
        i += 1

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


# ------------------------------------------------------------------
# Tabs
# ------------------------------------------------------------------
tab1, tab2 = st.tabs(["📊 Skill Gap & Roadmap Generator", "📄 AI Resume Scanner"])

# NOTE: the resume tab is written FIRST on purpose. Streamlit does not allow
# changing a widget's state after the widget is drawn, so the scanner must run
# before the checkboxes in tab1 are created. The visual tab order is unchanged.

# ---------------------------- TAB 2: Resume scanner ----------------------------
with tab2:
    st.subheader("Smart AI Resume Ingestion")
    st.write(
        "Upload your resume as PDF, Word (.docx), text, or an image (PNG/JPG). "
        "Detected skills will be auto-ticked in the first tab."
    )

    uploaded_file = st.file_uploader(
        "Upload Resume", type=["pdf", "docx", "txt"] + IMAGE_TYPES
    )

    if uploaded_file and st.button("Scan Resume with AI"):
        with st.spinner("Reading your resume and analyzing skills..."):
            all_skills = sorted({s for skills in roles_requirements.values() for s in skills})
            instructions = f"""You are a resume parser.
From the resume, return ONLY a JSON array of skills from this allowed list
that the candidate clearly has evidence of (projects, experience, skills section).
Do not add skills outside the list. No explanation, no markdown.

Allowed list: {json.dumps(all_skills)}"""

            name = uploaded_file.name.lower()
            reply = None

            if name.endswith(tuple(IMAGE_TYPES)):
                data = uploaded_file.getvalue()
                if len(data) > MAX_IMAGE_BYTES:
                    st.error("Image is too large. Please upload one under 4 MB.")
                else:
                    mime = "image/png" if name.endswith("png") else "image/jpeg"
                    reply = ask_llm(
                        instructions + "\n\nThe resume is in the attached image.",
                        temperature=0,
                        image_b64=base64.b64encode(data).decode(),
                        mime=mime,
                    )
            else:
                try:
                    resume_text = extract_resume_text(uploaded_file)
                except Exception as e:
                    st.error(f"Could not read this file: {e}")
                    resume_text = ""

                if not resume_text.strip():
                    st.error(
                        "Couldn't read text from this file (it may be a scanned image). "
                        "Try uploading it as a PNG/JPG instead."
                    )
                else:
                    reply = ask_llm(
                        instructions + "\n\nResume:\n" + resume_text[:MAX_RESUME_CHARS],
                        temperature=0,
                    )

            if reply is not None:
                found_lower = {f.lower() for f in parse_json_list(reply)}
                matched = []
                for role, skills in roles_requirements.items():
                    for skill in skills:
                        if skill.lower() in found_lower:
                            st.session_state[checkbox_key(role, skill)] = True
                            matched.append(skill)
                if matched:
                    st.success(
                        f"Detected {len(set(matched))} skills: "
                        f"{', '.join(sorted(set(matched)))}. "
                        "Switch to the first tab to see them ticked."
                    )
                else:
                    st.warning("No matching skills detected. Tick them manually in the first tab.")
    elif uploaded_file is None:
        st.info("Upload a file to begin.")

# ---------------------------- TAB 1: Skill gap ----------------------------
with tab1:
    col1, col2 = st.columns(2)

    with col1:
        st.subheader("1. Select Target Role")
        target_role = st.selectbox("Dream Career Role:", list(roles_requirements.keys()))
        required_skills = roles_requirements[target_role]

        st.subheader("2. Check Your Current Skills")
        user_skills = []
        for skill in required_skills:
            if st.checkbox(skill, key=checkbox_key(target_role, skill)):
                user_skills.append(skill)

    # Computed live, so results never disappear when other buttons are clicked
    missing_skills = [s for s in required_skills if s not in user_skills]
    match_percentage = int(len(user_skills) / len(required_skills) * 100)

    with col2:
        st.subheader("3. Results & Gap Analysis")
        st.metric(label="Role Readiness Match", value=f"{match_percentage}%")
        st.progress(match_percentage / 100)

        if missing_skills:
            st.warning(f"**Identified Skill Gaps:** {', '.join(missing_skills)}")
        else:
            st.success("Congratulations! You possess all core requirements listed for this role.")

    st.divider()

    # ---------------------------- AI roadmap ----------------------------
    st.subheader("4. AI Learning Roadmap")
    weeks = st.slider("Plan length (weeks)", 2, 12, 4)

    if st.button("Generate AI Learning Roadmap"):
        with st.spinner("Building your customized learning milestones..."):
            focus = ", ".join(missing_skills) if missing_skills else "None, focus on advanced portfolio projects"
            prompt = f"""Create a practical {weeks}-week structured learning plan and milestone checklist
for a student targeting the role of {target_role}.
Skills they already have: {', '.join(user_skills) if user_skills else 'None yet'}.
Their missing skills are: {focus}.
Format rules (important):
- Use plain markdown only: a "## Week N: <title>" heading for each week, then bullet points.
- Under each week use these bold labels: **Goal**, **Study topics**, **Free resources**, **Mini project**.
- Do NOT use tables. Do NOT use any HTML tags such as <br>.
- Put each topic, resource and project step on its own bullet line.
End with a "## Portfolio project idea" section that combines the skills. Keep it clear and concise."""
            text = ask_llm(prompt)
            if text:
                st.session_state["roadmap"] = {
                    "role": target_role,
                    "text": clean_roadmap_text(text),
                }

    roadmap = st.session_state.get("roadmap")
    if roadmap:
        st.subheader(f"🗺️ Custom Roadmap for {roadmap['role']}")
        st.markdown(roadmap["text"])
        safe_name = roadmap["role"].replace("/", "-").replace(" ", "_")
        st.download_button(
            "📄 Download roadmap (Word document)",
            data=markdown_to_docx(roadmap["text"], f"Learning Roadmap: {roadmap['role']}"),
            file_name=f"{safe_name}_roadmap.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )